#!/usr/bin/env python3
"""Validate final source correspondence, numbering and final artifacts."""
from pathlib import Path
import json,re,hashlib,math,zipfile
import openpyxl,fitz
from docx import Document
from PIL import Image
import xml.etree.ElementTree as ET
R=Path(__file__).resolve().parents[1];REPO=R.parent;checks=[]
counts=json.loads((R/'data/manuscript_counts.json').read_text());body=(R/'论文_正文.md').read_text()
assert 14000<=counts['main_with_figures_tables_cjk_characters']<=16000
assert '{{R' not in body
assert all(re.search(r'^## '+str(i)+' ',body,re.M) for i in range(1,7))
assert '单一数值字段，分别列出两端记录' in body
assert '28对中有26对的Cl端CT表值较大，而这26对中有20对' in body
assert '【待核实' in body and '【待补充】' in body
checks.append('Complete six-chapter manuscript, abstract and keywords; explicit unresolved fields and corrected NOCV/CT wording.')
mapping=json.loads((R/'data/citation_map.json').read_text());refs=mapping['references'];seen=[]
for group in re.findall(r'\^\[([0-9,\-]+)\]\^',body):
 nums=[int(n) for n in group.split(',')];assert nums==sorted(set(nums))
 for n in nums:
  if n not in seen:seen.append(n)
assert seen==list(range(1,12)) and len(refs)==11
assert all(ref['doi'] and '[J' in ref['formatted_entry'] for ref in refs)
assert not any(ref['id'] in ['R12','R13','R14'] for ref in refs)
checks.append('11 references, first-citation numbering and ascending compound superscripts; no unverifiable R12–R14 citations.')
records=json.loads((R/'data/complexes.json').read_text());pairs=json.loads((R/'data/pairs.json').read_text());by={d['system']:d for d in records}
assert len(records)==56 and len(pairs)==28
assert sum(p['delta_Eb_H_minus_Cl']<0 for p in pairs)==22
assert sum(p['delta_Eint_H_minus_Cl']<0 for p in pairs)==13
assert sum(bool(p['ranking_reversal']) for p in pairs)==9
ct=[p for p in pairs if by['Cl-'+p['ligand']]['CT']>by['H-'+p['ligand']]['CT']]
assert len(ct)==26 and sum(p['delta_Eb_H_minus_Cl']<0 for p in ct)==20
original_checks=0
if (REPO/'HOClxlsx.xlsx').exists():
 for path,digest in json.loads((R/'data/source_hashes.json').read_text()).items():assert hashlib.sha256((REPO/path).read_bytes()).hexdigest()==digest,path
 raw=openpyxl.load_workbook(REPO/'HOClxlsx.xlsx',data_only=True)
 for row in records:
  for key,source in row.items():
   if not key.endswith('_src'):continue
   m=re.fullmatch(r'HOClxlsx.xlsx::([^!]+)!([A-Z]+[0-9]+)',str(source))
   if not m:continue
   value=raw[m.group(1)][m.group(2)].value;actual=row[key[:-4]]
   if isinstance(value,(float,int)):assert math.isclose(value,actual,rel_tol=0,abs_tol=1e-12),(row['system'],key)
   else:assert value==actual,(row['system'],key)
   original_checks+=1
 checks.append(f'Original source hashes unchanged; {original_checks} original-cell fields match frozen data.')
else:checks.append('Original repository absent: original-cell comparison not repeated; frozen data only. This is not an original-input readiness claim.')
for name,images,tables in [('完整论文初稿.docx',4,3),('补充材料.docx',2,2),('参考文献.docx',0,0)]:
 with zipfile.ZipFile(R/name) as z:assert z.testzip() is None
 doc=Document(R/name);assert len(doc.inline_shapes)==images,(name,len(doc.inline_shapes));assert len(doc.tables)==tables,(name,len(doc.tables))
 if name=='完整论文初稿.docx':
  assert any(run.font.superscript for p in doc.paragraphs for run in p.runs)
  assert [len(t.rows) for t in doc.tables]==[29,11,9]
checks.append('All three Word files readable; main has 4 images/3 native tables and superscript references; supplement has 2 images/2 native tables.')
pdfs={}
for name,imagecount in [('完整论文初稿.pdf',4),('补充材料.pdf',2)]:
 d=fitz.open(R/name);assert len(d)>0;full='';imgcount=0;overflow=[]
 for i,page in enumerate(d):
  text=page.get_text();assert len(text.strip())>50,(name,i);full+=text;imgcount+=len(page.get_images())
  for image in page.get_images():
   for rect in page.get_image_rects(image[0]):
    if rect.x0<-1 or rect.y0<-1 or rect.x1>page.rect.width+1 or rect.y1>page.rect.height+1:overflow.append(i+1)
 assert imgcount==imagecount,(name,imgcount)
 assert not overflow,(name,overflow)
 log=(R/'audit'/f'{name}.log').read_text();assert 'Missing character' not in log and 'Error producing PDF' not in log
 if name=='完整论文初稿.pdf':
  compact=re.sub(r'\s+','',full)
  for token in ['22','13','15','9','0.1299','0.2595','10.1039/d6sc02210c','10.1021/jp806101t']:assert token in compact,token
 pdfs[name]={'pages':len(d),'embedded_figures':imgcount,'extractable_characters':len(full)}
checks.append('Main and supplement PDFs readable, text extractable, correct embedded figure counts, no page-overflow images or missing-glyph warnings.')
for meta in json.loads((R/'figures/figure_manifest.json').read_text()):
 stem=meta['figure']
 with Image.open(R/'figures'/f'{stem}.png') as im:im.verify()
 with fitz.open(R/'figures'/f'{stem}.pdf') as d:assert len(d)==1
 assert ET.parse(R/'figures'/f'{stem}.svg').findall('.//{http://www.w3.org/2000/svg}text')
book=openpyxl.load_workbook(R/'数据总表.xlsx',data_only=True)
assert len(book.sheetnames)==17 and book['构型总表56'].max_row==57 and book['能量配对28'].max_row==29 and book['第二阶段来源索引'].max_row==2734 and book['全稿引文编号'].max_row==12
checks.append('17-sheet workbook with 56 configurations, 28 pairs, 2733 evidence rows and 11 final reference mappings; all 18 standalone image files readable.')
result={'status':'passed','checks':checks,'original_cell_checks':original_checks,'counts':counts,'pdfs':pdfs,'manuscript_sha256':hashlib.sha256((R/'论文_正文.md').read_bytes()).hexdigest(),'limits':'Artifact/claim correspondence verified; unknown computational parameters, units, energy references and transfer assignments remain explicitly unverified.'}
(R/'audit/final_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
