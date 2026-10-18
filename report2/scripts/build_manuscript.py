#!/usr/bin/env python3
"""Build editable Word/Markdown, evidence tables and copied summary workbook."""
from pathlib import Path
import json,csv,re
import openpyxl
from openpyxl.styles import Font,PatternFill
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];REPO=R.parent
records=json.loads((R/'data/complexes.json').read_text());by={d['system']:d for d in records}
pairs=json.loads((R/'data/pairs.json').read_text());ids=json.loads((R/'data/identities.json').read_text())
caption={
'F1_energy_competition':('图3-5 28对体系的两端Eb与竞争差量（F1）','(a)表列H端与Cl端Eb配对；(b)δb=Eb(H)−Eb(Cl)；(c)δint=Eint(H)−Eint(Cl)。来源为HOClxlsx.xlsx的bess!B2:C29、F2:G29及编号对照。负差表示H端表列值较低，正差表示Cl端较低。阴影标记9对排序反转；(a)菱形标记三个H端候选及Cl-NHCme。(b)(c)采用标明的symlog坐标，±1以内线性，刻度仍为实际kcal/mol，不能按线段长度直接比较能量倍数。参考态及BSSE/ZPE未核实，图示不等于有限温度热力学或动力学偏好。'),
'F2_bond_response':('图3-4 HOCl内部键长响应（F2）','(a)全部28条H端Δr(O—H)；(b)其余25条H端记录的局部视图，三个优先候选仍完整保留在(a)，不将其余记录自动称为弱作用；(c)全部28条Cl端Δr(O—Cl)。差量由结构数据!E2:F57分别减去E59=0.97950、F59=1.71675复算。来源单位未注明，图轴保留原表距离单位；三H候选及Cl-NHCme单独标记。图示内部键响应，不能凭几何独自确定质子/氯完整转移。'),
'F3_eda_components':('图4-8 代表体系的完整EDA分项与总量（F3）','CO、N2CAACme、PPh3、NHCme和CAACmeSPh2各H/Cl两端，以ES、EX、REP、POL、ECORR原列名正负堆叠，菱形为表列EDA总量。来源为EDA!B5:G5/B33:G33、B8:G8/B36:G36、B23:G23/B51:G51、B14:G14/B42:G42、B4:G4/B32:G32。仅ES列头明确kcal/mol，本图其余分项暂按同单位展示并待核实；各子图使用独立且明示的纵轴尺度，不跨面板按柱高比较。保留正ECORR，不把它改称纯色散；全部分项参与抵消，不计算未定义贡献百分比。该图不能单独解释方向性或Eb排序。'),
'F4_geometry_candidates':('图5-1 新旧接触的分端几何筛查（F4）','(a)28条H端r(C—H)对r(O—H)；(b)28条Cl端r(C—Cl)对r(O—Cl)，来源结构数据!C2:F57。菱形突出H-CAACmeSPh2、H-NHCme、H-NHCmeSPh2及Cl-NHCme，方形为两端CO参照。距离保留原表数值，单位与原子号待核；不同面板比较不同原子对，不建立通用等距离转移阈值。此图为表列几何的真实数据图，不是优化分子结构图、NOCV图或反应路径；候选分类尚需新旧键配对电子结构证据。'),
'S1_energy_audit':('图S1 EDA求和与跨表能量审计','(a)R=EDA总量−(ES+EX+REP+POL+ECORR)；(b)EDA总量−bess Eint，均按体系名对应全部56条记录。阴影±0.03与±0.01分别为五项/总量和两列均独立保留0.01精度时的舍入包络，不是理论准确度或统计误差。H-N2NHCme残差+0.18、Cl-NHCmeSPh2跨表差+0.08、Cl-SCNHCme跨表差+0.50分别标注，未改任何原值。横轴标明symlog及线性区间±0.015；接近零不能证明方法、几何和参考态一致。单位限制同F3。'),
'S2_mep_competition':('图S2 同字段MEP与两端Eb差的探索性比较','横轴仅为MEP!B2:B29中24个非空的Vs,min(C)表值，表头单位kcal/mol；纵轴为bess所得δb，亦为kcal/mol，二者虽同单位却不是同一物理量。四个仅C列体系N2CS、SCCAACme、SCNHCme、SCPPh3未并入B列。(a)全部24点；(b)局部视图，三H端优先候选仍完整显示在(a)。图中蓝/橙表示Eb表内低值端，不代表已核实偏好；未作回归、未断言因果或单调关系。MEP表面、区域及计算条件待核实。')}
(R/'图注_第二阶段.md').write_text('# 第二阶段图注\n\n'+'\n\n'.join(t+'\n\n'+c for t,c in caption.values()))
def fmt(v,n=6):return '【待补充】' if v is None else (f'{v:.{n}f}' if isinstance(v,(int,float)) else str(v))
tables={}
tables['表3-1']=(['编号/体系','Eb(H)','Eb(Cl)','δb','δint','Eb表内低值端'],[[p['legacy_id']+' '+p['ligand'],fmt(p['H_Eb']),fmt(p['Cl_Eb']),fmt(p['delta_Eb_H_minus_Cl']),fmt(p['delta_Eint_H_minus_Cl']),p['Eb_numeric_lower']] for p in pairs],'能量单位kcal/mol；差量统一H减Cl；参考态/BSSE/ZPE待核实；全部来源bess列B/C/F/G，按编号对照逐行匹配。')
selected=['CO','N2CAACme','PPh3','NHCme','CAACmeSPh2'];erows=[]
for n in selected:
 for end in ['H','Cl']:
  d=by[end+'-'+n];erows.append([d['system']]+[fmt(d[k],2) for k in ['ES','EX','REP','POL','ECORR','EDA_E']])
tables['表4-3']=(['构型','ES','EX','REP','POL','ECORR','表列E'],erows,'按EDA原列名保留全部分项；仅ES明确标kcal/mol，其余同单位为审阅图假设。完整来源见D12和figure_F3 CSV。')
crow=[]
for n in ['H-CO','Cl-CO','H-NHCme','Cl-NHCme','H-CAACmeSPh2','Cl-CAACmeSPh2','H-NHCmeSPh2','Cl-NHCmeSPh2']:
 d=by[n];crow.append([n,fmt(d['r_CH'] if d['end']=='H' else d['r_CCl'],5),fmt(d['r_OH'] if d['end']=='H' else d['r_OCl'],5),fmt(d['rho_contact']),fmt(d['lap_contact']),fmt(d['E2max'],2),'H-CO已有，见正文' if n=='H-CO' else '【待补充】'])
tables['表5-1']=(['构型','C—H或C—Cl','O—H或O—Cl','表列ρ','表列∇²ρ','E(2)最大','新旧键WBI'],crow,'距离和拓扑字段单位/BCP原子对待核；E(2)明确kcal/mol。缺失WBI不填0；H-CAACmeSPh2的ρ差异需回查。源单元格见D03/D09/D10/D14/D15。')
groups={
'D01':('体系命名、覆盖与频率汇总',None,['min_frequency','imaginary_count']),
'D02':('MEP原字段及反例',None,[]),
'D03':('全部几何、基准与距离差',None,['r_CH','r_CCl','r_OH','r_OCl','dr_OH_cached','dr_OCl_cached','angle']),
'D04':('全28对能量排序及范围',None,['Eint_bess','Eb_bess']),
'D05':('九对反转与Eb-Eint表列差量',[p['ligand'] for p in pairs if p['ranking_reversal']],['Eint_bess','Eb_bess']),
'D06':('正或近零结合量',['CS','OCCS'],['Eint_bess','Eb_bess']),
'D07':('NBO通道身份、计数和E2',None,['donor','acceptor','E2max','E2sum_target','target_channel_count','cross_fragment_count','all_count']),
'D08':('CT端比较及与Eb反例',None,['CT','Eb_bess']),
'D09':('局部拓扑符号与rho可疑差异',None,['rho_contact','H_contact','lap_contact','r_CH','r_OH']),
'D10':('WBI可用与缺失',None,['WBI_CH','WBI_CCl','WBI_OH','WBI_OCl']),
'D11':('EDA完整求和与跨表差',None,['ES','EX','REP','POL','ECORR','EDA_E','Eint_bess']),
'D12':('代表体系EDA及配对抵消',selected,['ES','EX','REP','POL','ECORR','EDA_E','Eint_bess','Eb_bess']),
'D13':('NOCV Energy字段，含义未知',None,['NOCV_energy']),
'D14':('三个H端优先强重排候选',['CAACmeSPh2','NHCme','NHCmeSPh2'],['r_CH','r_OH','dr_OH_cached','rho_contact','H_contact','lap_contact','CT','E2max','E2sum_target','Eint_bess','Eb_bess']),
'D15':('Cl-NHCme重点个案',['NHCme'],['r_CCl','r_OCl','dr_OCl_cached','rho_contact','H_contact','lap_contact','CT','E2max','Eint_bess','Eb_bess']),
'D16':('按名称对应系列及混配比较',None,['Eint_bess','Eb_bess','r_CH','r_CCl','r_OH','r_OCl','CT']),
'D17':('近零端间差',['CAACme','PPh3','SPh2'],['Eint_bess','Eb_bess'])}
ledger=[]
for gid,(purpose,ligands,metrics) in groups.items():
 for d in records:
  if ligands is not None and d['ligand'] not in ligands:continue
  if gid=='D14' and d['end']!='H':continue
  if gid=='D15' and d['end']!='Cl':continue
  for k in metrics:ledger.append({'index':gid,'purpose':purpose,'system':d['system'],'metric':k,'value':d[k],'source':d[k+'_src'],'status':'原值；缺值为待补充，不能填0'})
 if gid in ['D01','D02']:
  for item in ids:
   if gid=='D01':ledger.append({'index':gid,'purpose':purpose,'system':item['ligand'],'metric':'identity','value':item['legacy_id']+'; '+item['H_system']+'; '+item['Cl_system'],'source':item['source'],'status':'仅名称映射；完整结构待核实'})
   else:
    for k in ['MEP_B','MEP_C']:ledger.append({'index':gid,'purpose':purpose,'system':item['ligand'],'metric':k,'value':item[k],'source':item[k+'_src'],'status':'B/C保持分栏；C单位和等值面未知'})
for gid in ['D03','D14','D15']:
 for cell,value in [('E59',0.9795),('F59',1.71675)]:ledger.append({'index':gid,'purpose':'表列HOCl基准','system':'HOCl','metric':cell,'value':value,'source':'HOClxlsx.xlsx::结构数据!'+cell,'status':'单位及与复合物同理论水平待核实'})
with (R/'data/数据来源索引.csv').open('w',encoding='utf-8-sig',newline='') as f:
 wr=csv.DictWriter(f,fieldnames=list(ledger[0]),lineterminator='\n');wr.writeheader();wr.writerows(ledger)
(R/'data/数据来源索引.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2))
refentries=[
'[1] LI S H, ZHOU H Y, LI Q Z, et al. Strong interactions between carbones and halogen atomic centers[J]. Chemical Science, 2026, 17: 14950-14962. DOI: 10.1039/d6sc02210c.',
'[2] ALKORTA I, BLANCO F, SOLIMANNEJAD M, et al. Competition of hydrogen bonds and halogen bonds in complexes of hypohalous acids with nitrogenated bases[J]. The Journal of Physical Chemistry A, 2008, 112: 10856-10863. DOI: 10.1021/jp806101t.',
'[3] HERRMANN B, SVATUNEK D. Directionality of halogen-bonds: insights from 2D energy decomposition analysis[J]. Chemistry—An Asian Journal, 2024, 19: e202301106. DOI: 10.1002/asia.202301106.',
'[4] ALKORTA I, LEGON A. Simple carbenes as hydrogen bond acceptors: ab initio determination of nucleophilicities and reduced nucleophilicities[J]. Physical Chemistry Chemical Physics, 2025, 27: 942-950. DOI: 10.1039/d4cp03894k.',
'[5] WIESKE L H E, ERDÉLYI M. Halogen bonds of halogen(I) ions—Where are we and where to go?[J]. Journal of the American Chemical Society, 2024, 146: 3-18. DOI: 10.1021/jacs.3c11449.']
(R/'参考文献_第二阶段.md').write_text('# 第二阶段实际引用文献\n\n'+'\n\n'.join(refentries)+'\n\n编号对应report1：1→R01，2→R02，3→R06，4→R05，5→R07。书目信息已在上传PDF内核对；未作在线二次核验。')
ledger_md='## 附录A 定量陈述来源表\n\n[Dxx]指向data/数据来源索引.csv；每条包含体系、指标、原值、原表单元格。下面给出用途和核心来源范围，逐值表为机器可复核版本。\n\n|索引|对应内容|原始来源|\n|---|---|---|\n'
for gid,(purpose,ligands,metrics) in groups.items():
 sheets=sorted({x['source'].split('::')[1].split('!')[0] for x in ledger if x['index']==gid})
 ledger_md+=f"|{gid}|{purpose}|HOClxlsx.xlsx：{'、'.join(sheets)}；逐值单元格见CSV|\n"
ledger_md+='\n衍生量统一为δb=Eb(H)−Eb(Cl)、δint=Eint(H)−Eint(Cl)、g=Eb−Eint、R=EDA_E−Σ5。D05表列差量未自动命名形变能；D11/S1的舍入界限仅用于算术审计。任何BCP、CT或NOCV字段身份未核实处均不得据文献反填。\n'
(R/'数据来源说明.md').write_text(ledger_md)
wb=openpyxl.load_workbook(REPO/'report1/统一数据总表_第一阶段.xlsx')
for sheetname in ['原始非空单元格','公式与缓存']:
 for row in wb[sheetname]:
  for cell in row:
   if cell.data_type=='f':cell.data_type='s'
ws=wb.create_sheet('第二阶段来源索引');keys=list(ledger[0]);ws.append(keys)
for row in ledger:ws.append([row.get(k) for k in keys])
for title,(heads,rows,note) in tables.items():
 ws=wb.create_sheet(title+'_第二阶段');ws.append(heads)
 for row in rows:ws.append(row)
 ws.append([note])
ws=wb.create_sheet('图注_第二阶段');ws.append(['文件前缀','图题','图注'])
for stem,(title,cap) in caption.items():ws.append([stem,title,cap])
for ws in wb:
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='305A72')
 if ws.title.endswith('第二阶段') or ws.title=='第二阶段来源索引':
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(70,max(16,max(len(str(c.value or '')) for c in col)+1))
wb.save(R/'统一数据总表_第二阶段.xlsx')
body=(R/'结果与讨论_第二阶段.md').read_text()
def mdtable(label):
 heads,rows,note=tables[label]
 return '\n\n'+label+'\n\n| '+' | '.join(heads)+' |\n| '+' | '.join('---' for x in heads)+' |\n'+'\n'.join('| '+' | '.join(row)+' |' for row in rows)+'\n\n'+note+'\n\n'
def mdfig(stem):
 title,cap=caption[stem];return f'\n\n![{title}](figures/{stem}.png)\n\n{title}。{cap}\n\n'
inserts={'### 3.3':mdfig('F2_bond_response'),'### 3.4':mdfig('F1_energy_competition')+mdtable('表3-1'),'### 4.4':mdfig('F3_eda_components')+mdtable('表4-3'),'### 5.2':mdfig('F4_geometry_candidates')+mdtable('表5-1')}
combined=[]
for line in body.splitlines():
 for key,value in inserts.items():
  if line.startswith(key):combined.append(value)
 combined.append(line)
full='\n'.join(combined)+'\n\n## 参考文献\n\n'+'\n\n'.join(refentries)+'\n\n'+ledger_md+'\n\n## 附录B 补充图\n\n'+mdfig('S1_energy_audit')+mdfig('S2_mep_competition')
(R/'结果与讨论_含图表.md').write_text(full)
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
for name in ['Normal','Body Text']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10.5);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.paragraph_format.line_spacing=1.35;st.paragraph_format.space_after=Pt(6)
for name,size in [('Title',17),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('183E51');st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
foot=sec.footer.paragraphs[0];foot.alignment=1;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def addtext(p,text):
 pos=0
 for m in re.finditer(r'\^\[([0-9,\-]+)\]\^',text):
  p.add_run(text[pos:m.start()]);run=p.add_run('['+m.group(1)+']');run.font.superscript=True;pos=m.end()
 p.add_run(text[pos:])
lines=full.splitlines();i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
 if m:doc.add_picture(str(R/m.group(2)),width=Cm(17.5));i+=1;continue
 if line.startswith('|') and i+1<len(lines) and lines[i+1].replace(' ','').startswith('|---'):
  cells=[x.strip() for x in line.strip('|').split('|')];table=doc.add_table(rows=1,cols=len(cells));table.style='Table Grid'
  for cell,text in zip(table.rows[0].cells,cells):cell.text=text
  hdr=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(hdr);i+=2
  while i<len(lines) and lines[i].lstrip().startswith('|'):
   row=table.add_row();vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   for cell,text in zip(row.cells,vals):cell.text=text
   i+=1
  for ri,row in enumerate(table.rows):
   for cell in row.cells:
    for p in cell.paragraphs:
     p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2)
     for run in p.runs:run.font.size=Pt(8);run.bold=(ri==0)
  continue
 if line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 elif line.startswith('### '):doc.add_heading(line[4:],2)
 else:
  p=doc.add_paragraph();addtext(p,line)
  if re.match(r'^图[0-9S]',line):
   p.paragraph_format.space_after=Pt(10)
   for run in p.runs:run.font.size=Pt(9)
 i+=1
doc.save(R/'结果与讨论_第二阶段.docx')
counts={'body_total_characters':len(body),'body_chinese_characters':len(re.findall(r'[\u4e00-\u9fff]',body)),'reference_count':5,'figures':6,'tables':3,'source_ledger_rows':len(ledger),'docx_paragraphs':len(doc.paragraphs),'docx_tables':len(doc.tables),'docx_images':len(doc.inline_shapes),'note':'字符统计包含标签/英文/数字；不声称已完成约1.5万字完整论文。'}
(R/'data/deliverable_counts.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2))
print(json.dumps(counts,ensure_ascii=False))
