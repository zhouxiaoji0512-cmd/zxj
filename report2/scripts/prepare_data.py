#!/usr/bin/env python3
"""Export stage-one workbook, preserving original values and source paths."""
from pathlib import Path
from docx import Document
import openpyxl,json,hashlib,csv
OUT=Path(__file__).resolve().parents[1];REPO=OUT.parent
SRC=REPO/'report1/统一数据总表_第一阶段.xlsx'
for folder in ['data','notes','figures']:(OUT/folder).mkdir(parents=True,exist_ok=True)
source_files=[SRC,SRC.parent/'第一阶段核查报告.docx',SRC.parent/'参考文献清单.docx',REPO/'HOClxlsx.xlsx',REPO/'论文大纲.docx']
hashes={str(p.relative_to(REPO)):hashlib.sha256(p.read_bytes()).hexdigest() for p in source_files}
manifest=OUT/'data/source_hashes.json'
if manifest.exists() and json.loads(manifest.read_text())!=hashes:raise RuntimeError('Source files changed; inspect before regenerating.')
w=openpyxl.load_workbook(SRC,data_only=True)
for sheet,name in [('构型总表56','complexes'),('能量配对28','pairs'),('体系对照','identities'),('指标字典','metric_dictionary'),('逐值来源','provenance'),('参考文献','references'),('图表计划','figure_plan')]:
 rows=list(w[sheet].values);heads=rows[0]
 data=[dict(zip(heads,row)) for row in rows[1:] if any(v is not None for v in row)]
 (OUT/'data'/f'{name}.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
 with (OUT/'data'/f'{name}.csv').open('w',encoding='utf-8-sig',newline='') as f:
  writer=csv.DictWriter(f,fieldnames=heads,lineterminator='\n');writer.writeheader();writer.writerows(data)
for n,fn in [('第一阶段核查报告.docx','report1_audit.txt'),('参考文献清单.docx','report1_references.txt')]:
 d=Document(SRC.parent/n);parts=[p.text for p in d.paragraphs if p.text]
 for i,t in enumerate(d.tables):
  parts.append('TABLE '+str(i+1));parts.extend(' | '.join(c.text for c in row.cells) for row in t.rows)
 (OUT/'notes'/fn).write_text('\n'.join(parts),encoding='utf-8')
manifest.write_text(json.dumps(hashes,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(json.loads((OUT/'data/complexes.json').read_text()))==56
assert len(json.loads((OUT/'data/pairs.json').read_text()))==28
print('56 configurations, 28 pairs exported; source hashes recorded.')
