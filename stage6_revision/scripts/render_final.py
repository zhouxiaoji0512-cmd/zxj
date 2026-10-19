#!/usr/bin/env python3
"""Render reviewed Markdown as native editable Word and read-only PDF."""
from pathlib import Path
import re,os,json,subprocess
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
def inline(p,text):
 text=text.replace('\\*','*').replace('\\_','_').replace('`','')
 pos=0
 for m in re.finditer(r'\^\[([0-9,\-]+)\]\^',text):
  p.add_run(text[pos:m.start()]);run=p.add_run('['+m.group(1)+']');run.font.superscript=True;pos=m.end()
 p.add_run(text[pos:])
def docx(source,target):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
 for name in ['Normal','Body Text']:
  style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(10.5);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');style.paragraph_format.line_spacing=1.35;style.paragraph_format.space_after=Pt(6)
 for name,size in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
  style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('173C52');style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 footer=sec.footer.paragraphs[0];footer.alignment=1;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 lines=source.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  if match:
   p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(R/match.group(2)),width=Cm(17.4));i+=1;continue
  if line.startswith('|') and i+1<len(lines) and lines[i+1].replace(' ','').startswith('|---'):
   headings=[x.strip() for x in line.strip('|').split('|')];table=doc.add_table(rows=1,cols=len(headings));table.style='Table Grid'
   for cell,value in zip(table.rows[0].cells,headings):cell.text=value
   repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat);i+=2
   while i<len(lines) and lines[i].lstrip().startswith('|'):
    vals=[x.strip() for x in lines[i].strip().strip('|').split('|')];row=table.add_row()
    for cell,value in zip(row.cells,vals):inline(cell.paragraphs[0],value)
    i+=1
   for ri,row in enumerate(table.rows):
    for cell in row.cells:
     for p in cell.paragraphs:
      p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(2)
      for run in p.runs:run.font.size=Pt(8);run.bold=(ri==0)
   continue
  if line.startswith('# '):doc.add_heading(line[2:],0)
  elif line.startswith('## '):doc.add_heading(line[3:],1)
  elif line.startswith('### '):doc.add_heading(line[4:],2)
  else:
   p=doc.add_paragraph();inline(p,line)
   if re.match(r'^图[0-9S]',line):
    p.paragraph_format.keep_together=True
    for run in p.runs:run.font.size=Pt(9)
   if line.startswith('【待核实'):
    for run in p.runs:run.font.color.rgb=RGBColor.from_string('9B3139')
  i+=1
 doc.save(target)
 return {'file':target.name,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'superscript_citations':sum(bool(run.font.superscript) for p in doc.paragraphs for run in p.runs)}
items=[('论文_图文整合.md','完整论文修订稿.docx','完整论文修订稿.pdf'),('补充材料.md','补充材料.docx','补充材料.pdf'),('参考文献.md','参考文献.docx',None),('作者确认表.md','作者确认表.docx',None),('本轮修订说明.md','本轮修订说明.docx',None)]
exports=[]
for src,word,pdf in items:exports.append(docx(R/src,R/word))
(R/'audit/word_exports.json').write_text(json.dumps(exports,ensure_ascii=False,indent=2))
print('Native Word files rendered.',flush=True)
env=os.environ.copy();env['XDG_CACHE_HOME']='/tmp/hocl-final-fontcache';Path(env['XDG_CACHE_HOME']).mkdir(exist_ok=True)
for src,word,pdf in items:
 if not pdf:continue
 text=(R/src).read_text()
 def together(match):
  path=match.group(1);cap=match.group(2)
  for oldchar,newtext in [('&',r'\&'),('%',r'\%'),('_',r'\_'),('#',r'\#')]:cap=cap.replace(oldchar,newtext)
  return '\n\n'+r'\begin{minipage}{\linewidth}'+ '\n'+r'\centering'+ '\n'+r'\includegraphics[width=\linewidth]{'+path+'}'+r'\par'+ '\n'+r'{\small '+cap+'}'+ '\n'+r'\end{minipage}'+ '\n\n'
 text=re.sub(r'!\[\]\(([^)]+)\)\{width=100%\}\s*\n\n([^\n]+)',together,text)
 submap=str.maketrans('₀₁₂₃₄₅₆₇₈₉','0123456789')
 supmap=str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻','0123456789+-')
 text=re.sub(r'[₀₁₂₃₄₅₆₇₈₉]+',lambda m:'$_{'+m.group().translate(submap)+'}$',text)
 text=re.sub(r'[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+',lambda m:'$^{'+m.group().translate(supmap)+'}$',text)
 pdfsource=R/'audit'/(src+'.pdf-source.md');pdfsource.write_text(text)
 cmd=['pandoc',str(pdfsource),'--resource-path',str(R),'--from=markdown+superscript','--pdf-engine=xelatex','-V','mainfont=Noto Serif CJK SC','-V','sansfont=Noto Sans CJK SC','-V','monofont=DejaVu Sans Mono','-V','fontsize=11pt','-V','geometry:margin=1.8cm','-V','linestretch=1.25','-H',str(R/'scripts/pdf_header.tex'),'-o',str(R/pdf)]
 proc=subprocess.run(cmd,cwd=R,capture_output=True,text=True,env=env,timeout=150)
 (R/'audit'/f'{pdf}.log').write_text(proc.stdout+'\n'+proc.stderr)
 if proc.returncode:raise RuntimeError(f'PDF export failed: {pdf}: {proc.stderr[-2000:]}')
 if 'Missing character' in proc.stderr:raise RuntimeError(f'Missing PDF glyph: {pdf}')
 print('PDF rendered:',pdf,flush=True)
