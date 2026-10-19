#!/usr/bin/env python3
"""Add implemented-figure mapping and precise metadata; retain historical plans."""
from pathlib import Path
import json,csv
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment

R=Path(__file__).resolve().parents[1]
dictionary=json.loads((R/'data/metric_dictionary.json').read_text())
units={
    'log_name':'作业或日志名称，文本标识，无能量单位',
    'donor':'供体轨道标签，文本归属，无能量单位',
    'acceptor':'受体轨道标签，文本归属，无能量单位',
    'target_channel_count':'目标通道计数，整数，无量纲；筛选规则待核实',
    'cross_fragment_count':'跨片段通道计数，整数，无量纲；片段及筛选规则待核实',
    'all_count':'汇总通道计数，整数，无量纲；统计范围待核实',
    'imaginary_count':'虚频个数，整数计数，无量纲；原始频率输出仍待核实',
    'log_line':'日志来源行号，位置索引，无物理单位',
    'original_NBO_check_status':'原始NBO核查状态，文本标识，无物理单位',
    'E2sum_target':'目标通道能量求和字段，单位和筛选规则待核实',
    'angle':'原表角度字段，单位及原子顺序待核实',
}
for key in ['r_CH','r_CCl','r_OH','r_OCl','dr_OH_cached','dr_OCl_cached']:
    units[key]='原表未标距离单位；数值保留，待原始几何输出确认'
for row in dictionary:
    if row['metric'] in units:row['unit_status']=units[row['metric']]
(R/'data/metric_dictionary.json').write_text(json.dumps(dictionary,ensure_ascii=False,indent=2))
with (R/'data/metric_dictionary.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(dictionary[0]));w.writeheader();w.writerows(dictionary)

captions=json.loads((R/'data/figure_captions.json').read_text())
definitions=[
('图1','F2','F2_bond_response','图3-4','三面板：H端全28、H端其余25局部、Cl端28','图1；3.2节','figure_F2_internal_bond_response.csv','内部键响应；单位待核实，非转移确证'),
('图2','F1','F1_energy_competition','图3-5','三面板：两端Eb、δb、δint；菱形候选，9对反转阴影','图2和表1；3.3节','figure_F1_energy_competition.csv','表内低值端及能量指标反转；非自由能选择性'),
('图3','F3','F3_eda_components','图4-8','五对体系、五分项正负堆叠及总量；独立纵轴','图3和表2；4.3节','figure_F3_EDA_components.csv','分项抵消；不等同纯色散、方向性或Eb排序'),
('图4','F4','F4_geometry_candidates','图5-1','两面板几何散点，各端28；四候选及CO参照','图4和表3；5.1节','figure_F4_geometry_screening.csv','新旧接触同步响应；不是优化结构或路径图'),
('图S1','S1','S1_energy_audit','原S1计划已替换','两面板：五项闭合残差、EDA与bess差；全56','补充材料S3；4.3节引用','figure_S1_energy_audit.csv','三处异常及舍入范围；不是理论误差'),
('图S2','S2','S2_mep_competition','原S2同字段MEP方案','两面板：B列24点全图及局部；C列4条不并入','补充材料S3；3.1节引用','figure_S2_MEP_exploratory.csv','同字段探索；不证明单调因果，非MEP表面'),
]
rows=[]
for number,asset,stem,old,design,location,data,limit in definitions:
    rows.append(dict(当前编号=number,内部资产标识=asset,文件主干=stem,历史编号或计划=old,
                     已实施设计=design,正文位置=location,绘图数据='data/'+data,
                     图题=captions[stem][0],证据范围=limit,输出格式='300 dpi PNG / SVG / PDF',状态='通用修订已完成；目标期刊体例待指定'))
(R/'data/implemented_figure_index.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
with (R/'data/implemented_figure_index.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

book=openpyxl.load_workbook(R/'数据总表.xlsx')
ws=book['指标字典']
for i,row in enumerate(dictionary,2):
    assert ws.cell(i,1).value==row['metric']
    ws.cell(i,3,row['unit_status'])
for name in ['图表计划','图注_第二阶段']:
    book[name]['A1'].comment=Comment('历史方案/第二阶段记录，保留原值。本轮实际成图及编号见“本轮图表索引”。','修订说明')
    book[name].sheet_properties.tabColor='999999'
for name in ['本轮图表索引','本轮修订说明']:
    if name in book:del book[name]
ws=book.create_sheet('本轮图表索引');ws.append(list(rows[0]))
for row in rows:ws.append(list(row.values()))
ws=book.create_sheet('本轮修订说明');ws.append(['项目','说明'])
counts=json.loads((R/'data/manuscript_counts.json').read_text())
for item in [
('日期','2026-10-19'),('基线','final_paper，main提交020f74d；基线目录保留'),
('当前状态','通用语言与图表衔接修订；用户选择指定期刊，期刊名称尚未提供'),
('数据','全部原始计算数据保持不变；修订仅涉及表述、图表索引和字段说明'),
('历史表','图表计划、图注_第二阶段及旧编号结果表保留历史信息，不能当作本轮最终编号'),
('字数口径',f"正文{counts['body_cjk_characters']}汉字，含正文图表说明{counts['main_with_figures_tables_cjk_characters']}汉字"),
('新图号','图1=F2；图2=F1；图3=F3；图4=F4；补图S1/S2保持'),
('新表号','表1=历史表3-1；表2=历史表4-3；表3=历史表5-1'),
('尚待核实','实际方法、完整体系身份、能量参考及校正、若干单位、EDA异常、候选成键分类'),
('期刊体例','摘要长度、章节转换、双语要素、图宽和参考文献细则须待目标期刊明确后处理'),
]:ws.append(item)
for name in ['本轮图表索引','本轮修订说明']:
    ws=book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='305A72')
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width=35
        for c in column:c.alignment=Alignment(wrap_text=True,vertical='top')
book.save(R/'数据总表.xlsx')
print('Updated metadata and implemented figure index;',len(book.sheetnames),'sheets. Numeric sheets untouched.')
